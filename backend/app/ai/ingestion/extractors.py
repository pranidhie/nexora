"""Safely extract text from approved NEXORA knowledge sources."""

from __future__ import annotations

import hashlib
import re
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.ai.ingestion.manifest import (
    KnowledgeManifest,
    SourceDocument,
)
from app.ai.ingestion.models import (
    ExtractedDocument,
    ExtractedSection,
)


MAX_SOURCE_FILE_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 500
MAX_WORKBOOK_SHEETS = 100
MAX_ROWS_PER_SHEET = 10_000
MAX_SECTION_CHARACTERS = 250_000

MARKDOWN_HEADING_PATTERN = re.compile(
    r"^(#{1,6})\s+(.+?)\s*$"
)


class DocumentExtractionError(ValueError):
    """Raised when an approved document cannot be safely extracted."""


def _normalise_text(value: str) -> str:
    """Normalise whitespace without destroying document structure."""

    normalised_lines = [
        line.rstrip()
        for line in value.replace(
            "\r\n",
            "\n",
        ).replace(
            "\r",
            "\n",
        ).split("\n")
    ]

    output_lines: list[str] = []
    previous_line_was_empty = False

    for line in normalised_lines:
        stripped_line = line.strip()

        if not stripped_line:
            if not previous_line_was_empty:
                output_lines.append("")

            previous_line_was_empty = True
            continue

        output_lines.append(stripped_line)
        previous_line_was_empty = False

    return "\n".join(output_lines).strip()


def _validate_source_file(
    source: SourceDocument,
) -> None:
    """Validate the physical file before opening it."""

    if not source.path.is_file():
        raise DocumentExtractionError(
            f"Approved source file was not found: "
            f"{source.relative_path}"
        )

    file_size = source.path.stat().st_size

    if file_size == 0:
        raise DocumentExtractionError(
            f"Approved source file is empty: "
            f"{source.relative_path}"
        )

    if file_size > MAX_SOURCE_FILE_BYTES:
        raise DocumentExtractionError(
            f"Approved source file exceeds the "
            f"{MAX_SOURCE_FILE_BYTES}-byte extraction limit: "
            f"{source.relative_path}"
        )


def _calculate_sha256(
    file_path: Path,
) -> str:
    """Calculate a reproducible SHA-256 checksum for a source file."""

    digest = hashlib.sha256()

    with file_path.open("rb") as source_file:
        for block in iter(
            lambda: source_file.read(65_536),
            b"",
        ):
            digest.update(block)

    return digest.hexdigest()


def _validate_section_text(
    source: SourceDocument,
    locator: str,
    text: str,
) -> str:
    """Validate and normalise extracted section text."""

    normalised_text = _normalise_text(text)

    if len(normalised_text) > MAX_SECTION_CHARACTERS:
        raise DocumentExtractionError(
            f"Extracted section is too large: "
            f"{source.source_id} at {locator}"
        )

    return normalised_text


def _extract_markdown_sections(
    source: SourceDocument,
) -> tuple[ExtractedSection, ...]:
    """Extract Markdown content while preserving heading provenance."""

    try:
        markdown_text = source.path.read_text(
            encoding="utf-8"
        )
    except UnicodeDecodeError as error:
        raise DocumentExtractionError(
            f"Markdown source is not valid UTF-8: "
            f"{source.relative_path}"
        ) from error
    except OSError as error:
        raise DocumentExtractionError(
            f"Markdown source could not be read: "
            f"{source.relative_path}"
        ) from error

    lines = markdown_text.splitlines()

    sections: list[ExtractedSection] = []
    current_title = source.title
    current_locator = "markdown:introduction"
    current_lines: list[str] = []
    section_number = 1

    def save_current_section() -> None:
        nonlocal section_number

        section_text = _validate_section_text(
            source,
            current_locator,
            "\n".join(current_lines),
        )

        if not section_text:
            return

        sections.append(
            ExtractedSection(
                section_id=(
                    f"{source.source_id}:"
                    f"section:{section_number}"
                ),
                source_id=source.source_id,
                title=current_title,
                locator=current_locator,
                text=section_text,
            )
        )

        section_number += 1

    for line_number, line in enumerate(
        lines,
        start=1,
    ):
        heading_match = MARKDOWN_HEADING_PATTERN.match(
            line
        )

        if heading_match:
            save_current_section()

            current_title = heading_match.group(2).strip()
            current_locator = (
                f"markdown:line:{line_number}"
            )
            current_lines = [current_title]
            continue

        current_lines.append(line)

    save_current_section()

    if not sections:
        raise DocumentExtractionError(
            f"No readable Markdown content was extracted from "
            f"{source.relative_path}"
        )

    return tuple(sections)


def _extract_pdf_sections(
    source: SourceDocument,
) -> tuple[ExtractedSection, ...]:
    """Extract one traceable section for each readable PDF page."""

    try:
        reader = PdfReader(
            str(source.path)
        )
    except (PdfReadError, OSError) as error:
        raise DocumentExtractionError(
            f"PDF source could not be opened: "
            f"{source.relative_path}"
        ) from error

    if reader.is_encrypted:
        raise DocumentExtractionError(
            f"Encrypted PDF sources are not supported: "
            f"{source.relative_path}"
        )

    if len(reader.pages) > MAX_PDF_PAGES:
        raise DocumentExtractionError(
            f"PDF exceeds the {MAX_PDF_PAGES}-page "
            f"extraction limit: {source.relative_path}"
        )

    sections: list[ExtractedSection] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            page_text = page.extract_text() or ""
        except Exception as error:
            raise DocumentExtractionError(
                f"PDF page extraction failed for "
                f"{source.source_id}, page {page_number}."
            ) from error

        locator = f"page:{page_number}"

        normalised_text = _validate_section_text(
            source,
            locator,
            page_text,
        )

        if not normalised_text:
            continue

        sections.append(
            ExtractedSection(
                section_id=(
                    f"{source.source_id}:"
                    f"page:{page_number}"
                ),
                source_id=source.source_id,
                title=f"Page {page_number}",
                locator=locator,
                text=normalised_text,
            )
        )

    if not sections:
        raise DocumentExtractionError(
            f"No readable PDF text was extracted from "
            f"{source.relative_path}"
        )

    return tuple(sections)


def _format_workbook_value(
    value: Any,
) -> str:
    """Convert a workbook cell value into deterministic text."""

    if value is None:
        return ""

    if isinstance(
        value,
        (datetime, date, time),
    ):
        return value.isoformat()

    if isinstance(value, Decimal):
        return format(value, "f")

    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"

    return str(value).replace(
        "\r",
        " ",
    ).replace(
        "\n",
        " ",
    ).strip()


def _extract_workbook_sections(
    source: SourceDocument,
) -> tuple[ExtractedSection, ...]:
    """Extract one traceable text section from each visible worksheet."""

    try:
        workbook = load_workbook(
            filename=source.path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as error:
        raise DocumentExtractionError(
            f"Excel source could not be opened: "
            f"{source.relative_path}"
        ) from error

    try:
        if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
            raise DocumentExtractionError(
                f"Workbook exceeds the "
                f"{MAX_WORKBOOK_SHEETS}-sheet extraction limit: "
                f"{source.relative_path}"
            )

        sections: list[ExtractedSection] = []

        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue

            extracted_rows: list[str] = []

            for row_number, row in enumerate(
                worksheet.iter_rows(
                    values_only=True
                ),
                start=1,
            ):
                if row_number > MAX_ROWS_PER_SHEET:
                    raise DocumentExtractionError(
                        f"Worksheet '{worksheet.title}' exceeds "
                        f"the {MAX_ROWS_PER_SHEET}-row limit."
                    )

                formatted_values = [
                    _format_workbook_value(value)
                    for value in row
                ]

                while (
                    formatted_values
                    and not formatted_values[-1]
                ):
                    formatted_values.pop()

                if not any(formatted_values):
                    continue

                extracted_rows.append(
                    "\t".join(formatted_values)
                )

            locator = (
                f"sheet:{worksheet.title}"
            )

            section_text = _validate_section_text(
                source,
                locator,
                "\n".join(extracted_rows),
            )

            if not section_text:
                continue

            sections.append(
                ExtractedSection(
                    section_id=(
                        f"{source.source_id}:"
                        f"sheet:{worksheet.title}"
                    ),
                    source_id=source.source_id,
                    title=worksheet.title,
                    locator=locator,
                    text=section_text,
                )
            )

        if not sections:
            raise DocumentExtractionError(
                f"No readable worksheet content was extracted "
                f"from {source.relative_path}"
            )

        return tuple(sections)
    finally:
        workbook.close()


def extract_source_document(
    source: SourceDocument,
) -> ExtractedDocument:
    """Extract one validated and approved knowledge source."""

    _validate_source_file(source)

    extractor_by_format = {
        "markdown": _extract_markdown_sections,
        "pdf": _extract_pdf_sections,
        "xlsx": _extract_workbook_sections,
    }

    extractor = extractor_by_format.get(
        source.format
    )

    if extractor is None:
        raise DocumentExtractionError(
            f"Unsupported extraction format: "
            f"{source.format}"
        )

    sections = extractor(source)

    return ExtractedDocument(
        source_id=source.source_id,
        title=source.title,
        relative_path=source.relative_path,
        format=source.format,
        document_type=source.document_type,
        phase=source.phase,
        domain=source.domain,
        classification=source.classification,
        tags=source.tags,
        file_sha256=_calculate_sha256(
            source.path
        ),
        sections=sections,
    )


def extract_manifest_documents(
    manifest: KnowledgeManifest,
) -> tuple[ExtractedDocument, ...]:
    """Extract all enabled sources from a validated manifest."""

    return tuple(
        extract_source_document(source)
        for source in manifest.sources
    )